"""Parse the FLDOE FAST K-2 Mathematics StudentData export (.xlsx).

FAST for grades K-2 is delivered on the Star platform, so each row carries a
"Unified Scale Score" (the Star Unified score), a "FAST Equivalent Score" (the
score on the FAST scale), a "Percentile Rank", and — crucially — the FAST
"Achievement Level" already computed. So we read the level straight from the
file (no concordance needed) and store it as a FAST Math result.

One row per student. The only identifier is the FLEID (e.g. "FL000012385738");
Local/District IDs come through as "N/A", so students are matched to the roster
by FLEID (which the roster import stores in flags), then by name.
"""
from __future__ import annotations

import io
import re

import openpyxl


def _s(v) -> str:
    return ("" if v is None else str(v)).strip()


def _num(v):
    v = _s(v).replace("%", "").replace(",", "")
    if v in ("", "N/A", "-"):
        return None
    try:
        return float(v)
    except ValueError:
        return None


def _level(v):
    m = re.search(r"(\d)", _s(v))
    return int(m.group(1)) if m else None


def _grade(v) -> str:
    g = _s(v).upper()
    if g.startswith("K") or g in ("0", "00", "KINDERGARTEN"):
        return "K"
    m = re.search(r"(\d)", g)   # "Grade 1" / "1" -> "1"
    return m.group(1) if m else g


def _period(reason: str) -> str:
    r = _s(reason).upper()
    for p in ("PM1", "PM2", "PM3"):
        if p in r.replace(" ", ""):
            return p
    return "PM1"


def is_fast_k2_export(headers) -> bool:
    joined = " ".join(_s(h).lower() for h in (headers or []))
    return ("fast mathematics" in joined or "fast reading" in joined
            or "fast ela" in joined) and "unified scale score" in joined \
        and "achievement level" in joined


def _headers(data: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(values_only=True), ())
    wb.close()
    return [_s(h) for h in header]


def detect(data: bytes) -> bool:
    try:
        return is_fast_k2_export(_headers(data))
    except Exception:
        return False


def parse_fast_k2(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        return {"students": [], "subject": "MATH", "period": "PM1"}
    hdr = [_s(h) for h in rows[0]]
    hl = [h.lower() for h in hdr]

    def find(*subs, exact=None):
        if exact:
            for j, h in enumerate(hl):
                if h == exact.lower():
                    return j
        for j, h in enumerate(hl):
            if all(s in h for s in subs):
                return j
        return None

    ci_id = find(exact="Student ID") or find("student id")
    ci_name = find("student name")
    ci_grade = find("enrolled grade")
    ci_reason = find("test reason")
    ci_uni = find("unified scale score")
    ci_equiv = find("fast equivalent score")
    ci_lvl = find("achievement level")
    ci_pct = find("percentile rank")
    ci_ell = find("english language")
    ci_ese = find("exceptionality")
    # Domain columns: "... Domain Score (out of 100)".
    dom_cols = [(j, re.sub(r"^\d+\.\s*", "", hdr[j]).replace(" Domain Score (out of 100)", "").strip())
                for j, h in enumerate(hl) if "domain score" in h]

    subject = "MATH" if "mathematics" in " ".join(hl) else "ELA"
    students = []
    period = "PM1"
    for r in rows[1:]:
        if ci_id is None or ci_id >= len(r):
            continue
        fleid = _s(r[ci_id])
        if not fleid:
            continue
        period = _period(r[ci_reason]) if ci_reason is not None else "PM1"
        name = _s(r[ci_name]) if ci_name is not None else ""
        last, first = "", ""
        if "," in name:
            last, first = [p.strip() for p in name.split(",", 1)]
        domains = {}
        for j, dn in dom_cols:
            val = _num(r[j]) if j < len(r) else None
            if val is not None:
                domains[dn] = val
        flags = {}
        if ci_ell is not None and _s(r[ci_ell]).upper() in ("Y", "YES", "LY", "LF"):
            flags["ell"] = _s(r[ci_ell])
        if ci_ese is not None and _s(r[ci_ese]) and _s(r[ci_ese]).upper() not in ("N/A", "NONE", ""):
            flags["ese"] = True
        students.append({
            "fleid": fleid,
            "first_name": first, "last_name": last,
            "grade": _grade(r[ci_grade]) if ci_grade is not None else "",
            "level": _level(r[ci_lvl]) if ci_lvl is not None else None,
            "scale_score": _num(r[ci_equiv]) if ci_equiv is not None else None,
            "unified": _num(r[ci_uni]) if ci_uni is not None else None,
            "percentile": _num(r[ci_pct]) if ci_pct is not None else None,
            "domains": domains, "flags": flags,
        })
    return {"students": students, "subject": subject, "period": period}
